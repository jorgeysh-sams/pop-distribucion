from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment


def generar_excel_resultado(resultado: dict, solicitud_id: int) -> BytesIO:
    """
    Recibe el dict que devuelve calcular_reparto() (con ok=True)
    y arma un Excel descargable en memoria.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Distribucion"

    # --- Encabezado con resumen de la solicitud ---
    ws["A1"] = "Solicitud No."
    ws["B1"] = solicitud_id
    ws["A2"] = "Tipo de POP"
    ws["B2"] = resultado.get("tipo_pop")
    ws["A3"] = "Modelo"
    ws["B3"] = resultado.get("modelo") or "(varios — ver detalle)"
    ws["A4"] = "Cantidad total"
    ws["B4"] = resultado.get("cantidad_total")
    ws["A5"] = "Total de registros"
    ws["B5"] = len(resultado.get("detalle", []))

    for row in range(1, 6):
        ws[f"A{row}"].font = Font(bold=True)

    # --- Tabla de detalle ---
    fila_inicio = 7
    headers = ["No.", "Modelo", "Tienda", "Account", "Site Group", "Cantidad Asignada"]
    header_fill = PatternFill(start_color="1F1F1F", end_color="1F1F1F", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")

    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=fila_inicio, column=col_idx, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    for i, item in enumerate(resultado.get("detalle", []), start=1):
        fila = fila_inicio + i
        ws.cell(row=fila, column=1, value=i)
        ws.cell(row=fila, column=2, value=item.get("modelo") or resultado.get("modelo"))
        ws.cell(row=fila, column=3, value=item.get("tienda"))
        ws.cell(row=fila, column=4, value=item.get("account"))
        ws.cell(row=fila, column=5, value=item.get("site_group"))
        ws.cell(row=fila, column=6, value=item.get("cantidad_asignada"))

    # --- Ajuste de anchos de columna ---
    anchos = [6, 30, 30, 25, 25, 20]
    for i, ancho in enumerate(anchos, start=1):
        ws.column_dimensions[chr(64 + i)].width = ancho

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
