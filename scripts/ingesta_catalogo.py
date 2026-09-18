"""
Carga un Excel con el catálogo de Material PoP hacia catalogo_pop.
Usa openpyxl puro (sin pandas/numpy) para evitar bloqueos de DLL en equipos
corporativos con políticas de Control de aplicaciones (WDAC/AppLocker).

Columnas esperadas en el Excel:
    Material PoP | Sku | Categoria de tienda | Division | Status

Uso:
    python scripts/ingesta_catalogo.py ruta_al_excel.xlsx [nombre_hoja]
"""
import os
import sys

from openpyxl import load_workbook
import psycopg2
from psycopg2.extras import execute_values
from dotenv import load_dotenv

load_dotenv()
DATABASE_URL = os.getenv("DATABASE_URL")

COLUMN_MAP = {
    "Material PoP": "material_pop",
    "Sku": "sku",
    "SKU": "sku",
    "Categoria de tienda": "categoria_tienda",
    "Categoría de tienda": "categoria_tienda",
    "Division": "division",
    "División": "division",
    "Status": "status",
}

COLUMNAS_REQUERIDAS = ["material_pop", "sku", "categoria_tienda", "division", "status"]


def cargar_catalogo(path: str, hoja=None):
    if not DATABASE_URL:
        raise RuntimeError("DATABASE_URL no está configurada. Revisa tu archivo .env")

    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb[hoja] if hoja else wb.worksheets[0]

    filas = ws.iter_rows(values_only=True)
    encabezados_originales = [str(h).strip() if h is not None else "" for h in next(filas)]
    encabezados_db = [COLUMN_MAP.get(h) for h in encabezados_originales]

    faltantes = [c for c in COLUMNAS_REQUERIDAS if c not in encabezados_db]
    if faltantes:
        raise ValueError(
            f"Faltan columnas en el Excel: {faltantes}. Encabezados encontrados: {encabezados_originales}"
        )

    valores = []
    for fila in filas:
        registro = {}
        for header_db, valor in zip(encabezados_db, fila):
            if header_db is not None:
                registro[header_db] = valor

        if not registro.get("material_pop") or not registro.get("sku"):
            continue

        status = registro.get("status") or "Activo"
        valores.append((
            registro.get("material_pop"),
            registro.get("sku"),
            registro.get("categoria_tienda"),
            registro.get("division"),
            status,
        ))

    conn = psycopg2.connect(DATABASE_URL)
    try:
        cur = conn.cursor()
        query = """
            INSERT INTO catalogo_pop (material_pop, sku, categoria_tienda, division, status)
            VALUES %s
            ON CONFLICT (material_pop, sku, division, categoria_tienda)
            DO UPDATE SET status = EXCLUDED.status
        """
        execute_values(cur, query, valores)
        conn.commit()
        print(f"✅ Procesadas {len(valores)} filas del catálogo.")
    finally:
        conn.close()


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("Uso: python scripts/ingesta_catalogo.py ruta_al_excel.xlsx [nombre_hoja]")
        sys.exit(1)

    ruta = sys.argv[1]
    hoja = sys.argv[2] if len(sys.argv) > 2 else None
    cargar_catalogo(ruta, hoja)
