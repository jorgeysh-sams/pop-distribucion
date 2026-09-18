"""
Carga un Excel de tiendas (RMF) hacia la tabla rmf_tiendas en Postgres.
Usa openpyxl puro (sin pandas/numpy) para evitar bloqueos de DLL en equipos
corporativos con políticas de Control de aplicaciones (WDAC/AppLocker).

Uso:
    python scripts/ingesta_excel.py ruta_al_excel.xlsx [nombre_hoja]
"""
import os
import sys

from openpyxl import load_workbook
import psycopg2
from psycopg2.extras import execute_values
from dotenv import load_dotenv

load_dotenv()
DATABASE_URL = os.getenv("DATABASE_URL")

COLUMN_MAP = {
    "GSCM": "gscm",
    "Country": "country",
    "City": "city",
    "Cluster Name": "cluster_name",
    "Sub Cluster": "sub_cluster",
    "Account": "account",
    "Site Group": "site_group",
    "MSO NAME": "mso_name",
    "MX": "mx",
    "AV": "av",
    "DA": "da",
    "GRADE POS": "grade_pos",
    "GRADE CLUSTER": "grade_cluster",
    "MX5": "mx5",
    "DA6": "da6",
    "VD": "vd",
    "EB": "eb",
    "Total7": "total7",
    "Coverage": "coverage",
}


def cargar_excel(path: str, hoja=None):
    if not DATABASE_URL:
        raise RuntimeError("DATABASE_URL no está configurada. Revisa tu archivo .env")

    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb[hoja] if hoja else wb.worksheets[0]

    filas = ws.iter_rows(values_only=True)
    encabezados_originales = [str(h).strip() if h is not None else "" for h in next(filas)]
    encabezados_db = [COLUMN_MAP.get(h) for h in encabezados_originales]

    columnas_validas = [c for c in encabezados_db if c is not None]
    if not columnas_validas:
        raise ValueError(
            "No se encontró ninguna columna esperada en el Excel. "
            "Revisa COLUMN_MAP en este script contra los encabezados reales: "
            f"{encabezados_originales}"
        )

    valores = []
    for fila in filas:
        registro = {}
        for header_db, valor in zip(encabezados_db, fila):
            if header_db is not None:
                registro[header_db] = valor

        if not any(registro.values()):
            continue  # fila vacía

        partes_desc = [f"{c}: {registro[c]}" for c in columnas_validas if registro.get(c) not in (None, "")]
        descripcion = " | ".join(partes_desc)

        valores.append(tuple(registro.get(c) for c in columnas_validas) + (descripcion,))

    columnas_insert = columnas_validas + ["descripcion_texto"]

    conn = psycopg2.connect(DATABASE_URL)
    try:
        cur = conn.cursor()
        query = f"INSERT INTO rmf_tiendas ({', '.join(columnas_insert)}) VALUES %s"
        execute_values(cur, query, valores)
        conn.commit()
        print(f"✅ Se insertaron {len(valores)} filas en rmf_tiendas.")
    finally:
        conn.close()


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("Uso: python scripts/ingesta_excel.py ruta_al_excel.xlsx [nombre_hoja]")
        sys.exit(1)

    ruta = sys.argv[1]
    hoja = sys.argv[2] if len(sys.argv) > 2 else None
    cargar_excel(ruta, hoja)
