"""
Carga inv_mso o inv_gsm desde un Excel.
Usa openpyxl puro (sin pandas/numpy) para evitar bloqueos de DLL en equipos
corporativos con políticas de Control de aplicaciones (WDAC/AppLocker).

Uso:
    python scripts/ingesta_inventarios.py inv_mso ruta_al_excel.xlsx [nombre_hoja]
    python scripts/ingesta_inventarios.py inv_gsm ruta_al_excel.xlsx [nombre_hoja]
"""
import os
import sys

from openpyxl import load_workbook
import psycopg2
from psycopg2.extras import execute_values
from dotenv import load_dotenv

load_dotenv()
DATABASE_URL = os.getenv("DATABASE_URL")

CONFIG = {
    "inv_mso": {
        "column_map": {
            "Combinada": "combinada",
            "Site Id": "site_id",
            "Modelo": "modelo",
            "Inventario": "inventario",
            "Prom Inv": "prom_inv",
            "Lineas": "lineas",
            "Recuento": "recuento",
        },
        "tabla": "inv_mso",
    },
    "inv_gsm": {
        "column_map": {
            "Site ID": "site_id",
            "MKT_Name": "mkt_name",
            "Inventario": "inventario",
            "Prom Inv": "prom_inv",
            "Lineas": "lineas",
        },
        "tabla": "inv_gsm",
    },
}


def cargar(tabla: str, path: str, hoja=None):
    if tabla not in CONFIG:
        raise ValueError(f"Tabla '{tabla}' no soportada. Usa: {list(CONFIG.keys())}")
    if not DATABASE_URL:
        raise RuntimeError("DATABASE_URL no está configurada. Revisa tu archivo .env")

    cfg = CONFIG[tabla]
    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb[hoja] if hoja else wb.worksheets[0]

    filas = ws.iter_rows(values_only=True)
    encabezados_originales = [str(h).strip() if h is not None else "" for h in next(filas)]
    encabezados_db = [cfg["column_map"].get(h) for h in encabezados_originales]

    columnas_validas = [c for c in encabezados_db if c is not None]
    if not columnas_validas:
        raise ValueError(
            f"No se encontró ninguna columna esperada para {tabla}. "
            f"Revisa CONFIG en este script contra los encabezados reales: {encabezados_originales}"
        )

    valores = []
    for fila in filas:
        registro = {}
        for header_db, valor in zip(encabezados_db, fila):
            if header_db is not None:
                registro[header_db] = valor

        if not any(registro.values()):
            continue

        partes_desc = [f"{c}: {registro[c]}" for c in columnas_validas if registro.get(c) not in (None, "")]
        descripcion = " | ".join(partes_desc)

        valores.append(tuple(registro.get(c) for c in columnas_validas) + (descripcion,))

    columnas_insert = columnas_validas + ["descripcion_texto"]

    conn = psycopg2.connect(DATABASE_URL)
    try:
        cur = conn.cursor()
        query = f"INSERT INTO {cfg['tabla']} ({', '.join(columnas_insert)}) VALUES %s"
        execute_values(cur, query, valores)
        conn.commit()
        print(f"✅ Se insertaron {len(valores)} filas en {cfg['tabla']}.")
    finally:
        conn.close()


if __name__ == "__main__":
    if len(sys.argv) < 3:
        print("Uso: python scripts/ingesta_inventarios.py [inv_mso|inv_gsm] ruta_al_excel.xlsx [nombre_hoja]")
        sys.exit(1)

    tabla = sys.argv[1]
    ruta = sys.argv[2]
    hoja = sys.argv[3] if len(sys.argv) > 3 else None
    cargar(tabla, ruta, hoja)
